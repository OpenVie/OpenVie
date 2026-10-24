# Copyright 2026 LinkedNodeDigital
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

from __future__ import annotations

import csv
import hashlib
import logging
import re
from abc import ABC, abstractmethod
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from decimal import Decimal
from html.parser import HTMLParser
from io import BytesIO, StringIO
from pathlib import PurePosixPath
from typing import Any, Literal, NoReturn

from app.modules.ingestion.api import PermanentIngestionFailure, TransientIngestionFailure
from app.modules.ingestion.internal.ocr import OcrIngestionService
from app.modules.ingestion.internal.transcription import AsrIngestionService, TranscriptResult

logger = logging.getLogger(__name__)

BlockType = Literal[
    "heading", "paragraph", "list", "code", "quote", "table", "page", "sheet", "row"
]

# Document-wide derived-text provenance. ``asr`` marks transcripts from
# optional audio/video ingestion; ``ocr`` marks recognized raster content.
TextSource = Literal["text", "ocr", "mixed", "asr"]

# Embedded-image OCR (best-effort) guards. Disabled by default; see
# ``INGESTION_OCR_EMBEDDED_IMAGES`` in app/bootstrap/settings.py.
_EMBEDDED_MIN_PIXELS = 32  # icons, bullets, and decorative rules carry no text
_EMBEDDED_MAX_IMAGES = 64  # per-document budget; worst-case OCR cost is bounded
_EMBEDDED_OCR_SOURCE = "embedded-image"
_EMBEDDED_PARAGRAPH_LIMIT = 8  # paragraphs kept per embedded image
# Provider-level failures are tolerated per image (best effort), but a streak
# this long means systematic breakage (missing tesseract, VLM outage, all
# timeouts): fail the document instead of silently dropping every image.
_EMBEDDED_PROVIDER_ERROR_LIMIT = 3

_IMAGE_MIME_BY_FORMAT = {
    "PNG": "image/png",
    "JPEG": "image/jpeg",
    "GIF": "image/gif",
    "BMP": "image/bmp",
    "WEBP": "image/webp",
    "TIFF": "image/tiff",
}

# Metadata keys bound to OCR-derived image blocks (round-trips through the
# OCR artifact JSON; bytes never reach it — assets are keyed by name).
_IMAGE_NAME_KEY = "image_name"


@dataclass(frozen=True, slots=True)
class ImageAsset:
    """A persisted raster image derived from a document.

    ``name`` identifies the asset inside its source document (media part name
    in DOCX/XLSX, pypdf image name in PDF, ``"source"`` for standalone image
    files) and gives the deterministic storage-key suffix. ``blob`` is empty
    for descriptors rebuilt from the OCR artifact cache: bytes are uploaded
    once at artifact-save time and never flow through cache round-trips.
    """

    name: str
    media_type: str
    width: int | None
    height: int | None
    blob: bytes = b""


def _image_asset(
    name: str, blob: bytes, *, min_pixels: int = _EMBEDDED_MIN_PIXELS
) -> ImageAsset | None:
    """Describe a decodable raster image; ``None`` when PIL cannot read it."""
    try:
        from PIL import Image
    except ImportError:  # pragma: no cover - dependency ships with the service
        return None
    try:
        with Image.open(BytesIO(blob)) as frame:
            width, height = frame.size
            categories = _IMAGE_MIME_BY_FORMAT.get((frame.format or "").upper())
    except Exception:  # noqa: BLE001 - best-effort probing of arbitrary media
        return None
    if width < min_pixels or height < min_pixels:
        return None
    if not categories:
        return None
    return ImageAsset(
        name=name,
        media_type=categories,
        width=int(width),
        height=int(height),
        blob=blob,
    )


def _image_block_metadata(asset: ImageAsset) -> dict[str, Any]:
    """Block metadata binding OCR text to its source image."""
    return {
        _IMAGE_NAME_KEY: asset.name,
        "image_media_type": asset.media_type,
        "image_width": asset.width,
        "image_height": asset.height,
    }


@dataclass(frozen=True, slots=True)
class ExtractedPage:
    """Compatibility view retained for callers that still consume page text."""

    page_number: int
    text: str


@dataclass(frozen=True, slots=True)
class KnowledgeBlock:
    unit_id: str
    block_type: BlockType
    text: str
    section_path: tuple[str, ...] = ()
    heading_context: str | None = None
    page_number: int | None = None
    sheet_name: str | None = None
    cell_range: str | None = None
    table_id: str | None = None
    source_start: int | None = None
    source_end: int | None = None
    content_hash: str = ""
    parser_version: str = "digital-v1"
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    modality: Literal["document", "spreadsheet"]
    blocks: tuple[KnowledgeBlock, ...]
    tables: tuple[NormalizedTable, ...] = ()
    text_source: TextSource = "text"
    page_count: int | None = None
    # Raster assets worth persisting as citation images (descriptors carry
    # blobs only on the live parse path; artifact caches drop them).
    media: tuple[ImageAsset, ...] = ()


@dataclass(frozen=True, slots=True)
class NormalizedColumn:
    name: str
    inferred_type: Literal["integer", "decimal", "date", "datetime", "boolean", "string"]


@dataclass(frozen=True, slots=True)
class NormalizedRow:
    row_number: int
    cell_range: str
    values: dict[str, Any]
    formula_columns: tuple[str, ...] = ()
    formula_expressions: dict[str, str] = field(default_factory=dict)


@dataclass(frozen=True, slots=True)
class NormalizedTable:
    table_id: str
    sheet_name: str
    cell_range: str
    columns: tuple[NormalizedColumn, ...]
    rows: tuple[NormalizedRow, ...]


def _block(
    source: str,
    ordinal: int,
    block_type: BlockType,
    text: str,
    **kwargs: Any,
) -> KnowledgeBlock:
    normalized = text.strip()
    identity = (
        f"{source}:{ordinal}:{block_type}:{kwargs.get('page_number')}:{kwargs.get('sheet_name')}"
    )
    return KnowledgeBlock(
        unit_id=hashlib.sha256(identity.encode()).hexdigest()[:32],
        block_type=block_type,
        text=normalized,
        content_hash=hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
        **kwargs,
    )


def _ocr_paragraphs(text: str, limit: int | None = None) -> tuple[str, ...]:
    """Split OCR text into non-empty paragraphs, optionally capped."""
    paragraphs = tuple(
        paragraph.strip() for paragraph in re.split(r"\n\s*\n", text) if paragraph.strip()
    )
    return paragraphs if limit is None else paragraphs[:limit]


def _fail_embedded_streak(file_name: str, count: int, cause: Exception) -> NoReturn:
    """Abort after repeated provider failures, preserving the original error kind.

    A streak of consecutive provider errors means systematic breakage (missing
    tesseract binary, VLM outage, every call timing out) — silently ingesting a
    document without any of its embedded images would hide the problem.
    """
    message = f"Embedded-image OCR failed {count} provider calls in a row for {file_name}"
    if isinstance(cause, TransientIngestionFailure):
        raise TransientIngestionFailure(message) from cause
    raise PermanentIngestionFailure(message) from cause


def _media_sort_key(name: str) -> tuple[str, ...]:
    """Natural sort so ``image2`` precedes ``image10`` like document order."""
    return tuple(
        chunk.zfill(8) if chunk.isdigit() else chunk for chunk in re.split(r"(\d+)", name)
    )


def _media_entries(
    data: bytes, directories: tuple[str, ...], label: str
) -> list[tuple[str, bytes]]:
    """Best-effort raster images embedded in an OOXML package.

    Only decodable images with both dimensions >= ``_EMBEDDED_MIN_PIXELS`` are
    returned, naturally sorted, capped at ``_EMBEDDED_MAX_IMAGES`` entries.
    """
    import zipfile

    try:
        from PIL import Image
    except ImportError as exc:  # pragma: no cover - dependency ships with the service
        raise PermanentIngestionFailure("Image decoder dependency is not installed") from exc

    entries: list[tuple[str, bytes]] = []
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            names = [name for name in archive.namelist() if name.startswith(directories)]
            for name in sorted(names, key=_media_sort_key):
                blob = archive.read(name)
                try:
                    with Image.open(BytesIO(blob)) as frame:
                        width, height = frame.size
                except Exception:
                    logger.warning(
                        "Skipping unreadable embedded image %s in %s", name, label
                    )
                    continue
                if width < _EMBEDDED_MIN_PIXELS or height < _EMBEDDED_MIN_PIXELS:
                    continue
                entries.append((name, blob))
    except Exception as exc:
        logger.warning("Unable to scan media parts of %s: %s", label, exc)
    return entries[:_EMBEDDED_MAX_IMAGES]


def _xlsx_media_sheets(data: bytes, label: str) -> dict[str, str]:
    """Map ``xl/media/*`` zip entries to the sheet title they are anchored on.

    Workbook and drawing relationship chains are read from the package; any
    structural surprise only shrinks the result (blocks keep ``sheet_name=None``).
    """
    import xml.etree.ElementTree as ET
    import zipfile

    main_ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    office_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

    def resolve(base: str, target: str) -> str:
        if target.startswith("/"):
            return target.lstrip("/")
        parts = [chunk for chunk in base.split("/") if chunk]
        for chunk in target.split("/"):
            if chunk in ("", "."):
                continue
            if chunk == "..":
                if parts:
                    parts.pop()
            else:
                parts.append(chunk)
        return "/".join(parts)

    def relationships(archive: Any, part: str) -> dict[str, str]:
        directory = part.rsplit("/", 1)[0] if "/" in part else "."
        base = part.rsplit("/", 1)[-1]
        rels = ET.fromstring(archive.read(f"{directory}/_rels/{base}.rels"))
        return {
            str(item.get("Id")): str(item.get("Target", ""))
            for item in rels.iter(f"{rel_ns}Relationship")
        }

    mapping: dict[str, str] = {}
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            wb_rels = relationships(archive, "xl/workbook.xml")
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            for sheet in workbook.iter(f"{main_ns}sheet"):
                raw_target = wb_rels.get(str(sheet.get(f"{office_ns}id") or "")) or ""
                sheet_path = raw_target.lstrip("/")
                if not sheet_path:
                    continue
                sheet_dir = sheet_path.rsplit("/", 1)[0] if "/" in sheet_path else "."
                sheet_rels = relationships(archive, sheet_path)
                for drawing_target in sheet_rels.values():
                    drawing_path = resolve(sheet_dir, drawing_target)
                    if not drawing_path.startswith("xl/drawings/drawing"):
                        continue
                    drawing_dir = drawing_path.rsplit("/", 1)[0]
                    for media_target in relationships(archive, drawing_path).values():
                        media_name = resolve(drawing_dir, media_target)
                        if media_name.startswith("xl/media/"):
                            title = sheet.get("name")
                            if title:
                                mapping[media_name] = str(title)
    except Exception as exc:  # noqa: BLE001 - provenance stays best-effort
        logger.warning("Unable to map embedded images to sheets of %s: %s", label, exc)
    return mapping


def _docx_inline_media_name(document: Any, rid: str) -> str | None:
    """Zip entry name of an inline picture relationship; None when unusable."""
    try:
        relationship = document.part.rels[rid]
    except KeyError:
        return None
    if getattr(relationship, "is_external", False):
        return None
    try:
        return str(relationship.target_part.partname).lstrip("/")
    except Exception:  # noqa: BLE001 - best-effort provenance
        return None


def _interleave_media_blocks(
    base: list[KnowledgeBlock],
    anchors: dict[str, tuple[int, tuple[str, ...], int]],
    media_plan: dict[str, list[KnowledgeBlock]],
) -> list[KnowledgeBlock]:
    """Place anchored media blocks at their image position; the rest trail."""
    placements: list[tuple[int, int, str]] = []
    trailing: list[KnowledgeBlock] = []
    for index, (name, media_blocks) in enumerate(media_plan.items()):
        if media_blocks and name in anchors:
            placements.append((anchors[name][2], index, name))
        else:
            trailing.extend(media_blocks)
    result = list(base)
    for position, _index, name in sorted(
        placements, key=lambda item: (-item[0], item[1])
    ):
        result[position:position] = media_plan[name]
    result.extend(trailing)
    return result


def _embedded_ocr_batches(
    ocr: OcrIngestionService,
    media: list[tuple[str, bytes]],
    file_name: str,
    container: str,
    *,
    start: int,
    page_number: int | None,
    anchors: dict[str, tuple[int, tuple[str, ...], int]] | None = None,
    sheets: dict[str, str] | None = None,
) -> tuple[list[tuple[str, list[KnowledgeBlock]]], tuple[ImageAsset, ...]]:
    """OCR embedded images best-effort, grouped per media with provenance.

    DOCX anchors give ``page_number`` + ``section_path`` (estimated page from
    page-break markers, heading context at the image position); XLSX sheets map
    images to their sheet title. Whatever fails is logged and skipped, and only
    a repeated provider failure raised by the service fails the document.
    Blocks are bound to their source image through metadata, and every image
    that actually produced text is returned as a citable asset.
    """
    batches: list[tuple[str, list[KnowledgeBlock]]] = []
    assets: list[ImageAsset] = []
    ordinal = start
    streak = 0
    for name, blob in media:
        try:
            recognition = ocr.ocr_embedded_image_best_effort(
                blob, content_type="", label=f"{file_name}#{name}"
            )
        except Exception as exc:  # noqa: BLE001 - best-effort with a streak cap
            streak += 1
            logger.warning("Skipping embedded image %s in %s: %s", name, file_name, exc)
            if streak >= _EMBEDDED_PROVIDER_ERROR_LIMIT:
                _fail_embedded_streak(file_name, streak, exc)
            continue
        streak = 0
        if recognition is None:
            continue
        asset = _image_asset(name, blob)
        metadata: dict[str, Any] = {
            "ocr_source": _EMBEDDED_OCR_SOURCE,
            "ocr_container": container,
        }
        if asset is not None:
            assets.append(asset)
            metadata.update(_image_block_metadata(asset))
        anchor = (anchors or {}).get(name)
        kwargs: dict[str, Any] = {"page_number": page_number}
        if anchor:
            (anchor_page, section_path, _position) = anchor
            kwargs["page_number"] = anchor_page
            kwargs["section_path"] = section_path
            kwargs["heading_context"] = section_path[-1] if section_path else None
        sheet = (sheets or {}).get(name)
        if sheet:
            kwargs["sheet_name"] = sheet
        grouped: list[KnowledgeBlock] = []
        for paragraph in _ocr_paragraphs(recognition.text, _EMBEDDED_PARAGRAPH_LIMIT):
            grouped.append(
                _block(
                    file_name,
                    ordinal,
                    "paragraph",
                    paragraph,
                    parser_version="ocr-v1",
                    metadata=metadata,
                    **kwargs,
                )
            )
            ordinal += 1
        batches.append((name, grouped))
    return batches, tuple(assets)


def _decode_utf8(data: bytes, label: str) -> str:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise PermanentIngestionFailure(f"{label} document is not valid UTF-8") from exc
    if not text.strip():
        raise PermanentIngestionFailure("Document contains no extractable text")
    return text


class ParserAdapter(ABC):
    extensions: tuple[str, ...] = ()
    content_types: tuple[str, ...] = ()

    @abstractmethod
    def parse(
        self, data: bytes, *, file_name: str, content_type: str = ""
    ) -> ParsedDocument: ...


class TextParser(ParserAdapter):
    extensions = (".txt",)
    content_types = ("text/plain",)

    def parse(
        self, data: bytes, *, file_name: str, content_type: str = ""
    ) -> ParsedDocument:
        text = _decode_utf8(data, "TXT")
        blocks: list[KnowledgeBlock] = []
        offset = 0
        for ordinal, match in enumerate(re.finditer(r"\S(?:.*?\S)?(?=\n\s*\n|\Z)", text, re.S)):
            paragraph = match.group(0).strip()
            blocks.append(
                _block(
                    file_name,
                    ordinal,
                    "paragraph",
                    paragraph,
                    page_number=1,
                    source_start=match.start(),
                    source_end=match.end(),
                )
            )
            offset = match.end()
        if not blocks and text.strip():
            blocks.append(
                _block(
                    file_name,
                    0,
                    "paragraph",
                    text,
                    page_number=1,
                    source_start=0,
                    source_end=len(text),
                )
            )
        del offset
        return ParsedDocument("document", tuple(blocks))


class PdfParser(ParserAdapter):
    """Digital PDF parser with an optional OCR fallback for scanned pages.

    A page whose native text layer yields fewer than ``min_text_chars`` visible
    characters is rasterized and recognized through the configured OCR service;
    everything else keeps the existing fast path byte-for-byte.
    """

    extensions = (".pdf",)
    content_types = ("application/pdf",)

    def __init__(self, ocr: OcrIngestionService | None = None):
        self._ocr = ocr

    def parse(
        self, data: bytes, *, file_name: str, content_type: str = ""
    ) -> ParsedDocument:
        if not data.startswith(b"%PDF-"):
            raise PermanentIngestionFailure("PDF document signature is invalid")
        try:
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise PermanentIngestionFailure("Encrypted PDF files are not supported")
        except PermanentIngestionFailure:
            raise
        except Exception as exc:
            raise PermanentIngestionFailure("PDF document could not be parsed") from exc
        page_count = len(reader.pages)
        native: dict[int, tuple[str, ...]] = {}
        ocr_pages: list[int] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            paragraphs = tuple(
                paragraph.strip()
                for paragraph in re.split(r"\n\s*\n", text)
                if paragraph.strip()
            )
            native[page_number] = paragraphs
            if self._ocr is not None and _visible_chars(paragraphs) < self._ocr.min_text_chars:
                ocr_pages.append(page_number)
        if ocr_pages:
            if self._ocr is None:
                raise PermanentIngestionFailure(
                    "PDF document contains no extractable text; scanned PDFs are not supported"
                )
            self._ocr.check_page_budget(page_count, label=file_name)
            recognized = self._ocr.ocr_pdf_pages(
                data, tuple(n - 1 for n in ocr_pages), label=file_name
            )
        else:
            recognized = {}
        embedded: dict[int, list[tuple[str, ImageAsset | None]]] = {}
        embedded_assets: tuple[ImageAsset, ...] = ()
        if self._ocr is not None and self._ocr.embedded_images:
            embedded, embedded_assets = self._embedded_pdf_text(
                reader,
                [page for page in range(1, page_count + 1) if page not in ocr_pages],
                label=file_name,
            )
        blocks: list[KnowledgeBlock] = []
        for page_number in range(1, page_count + 1):
            for paragraph in native[page_number]:
                blocks.append(
                    _block(
                        file_name,
                        len(blocks),
                        "page",
                        paragraph,
                        page_number=page_number,
                    )
                )
            if page_number in recognized:
                for paragraph in _ocr_paragraphs(recognized[page_number].text):
                    blocks.append(
                        _block(
                            file_name,
                            len(blocks),
                            "page",
                            paragraph,
                            page_number=page_number,
                            parser_version="ocr-v1",
                        )
                    )
            for paragraph, asset in embedded.get(page_number, ()):
                blocks.append(
                    _block(
                        file_name,
                        len(blocks),
                        "page",
                        paragraph,
                        page_number=page_number,
                        parser_version="ocr-v1",
                        metadata={
                            "ocr_source": _EMBEDDED_OCR_SOURCE,
                            **(_image_block_metadata(asset) if asset is not None else {}),
                        },
                    )
                )
        if not blocks:
            if self._ocr is not None:
                raise PermanentIngestionFailure(
                    "PDF document produced no text through the native text layer or OCR"
                )
            raise PermanentIngestionFailure(
                "PDF document contains no extractable text; scanned PDFs are not supported"
            )
        return ParsedDocument(
            "document",
            tuple(blocks),
            text_source=_text_source(native, sorted((*ocr_pages, *embedded)), page_count),
            page_count=page_count,
            media=tuple(embedded_assets),
        )

    def _embedded_pdf_text(
        self, reader: Any, pages: list[int], *, label: str
    ) -> tuple[dict[int, list[tuple[str, ImageAsset | None]]], tuple[ImageAsset, ...]]:
        """OCR images embedded in text-bearing PDF pages (best effort).

        Pages that failed the text-layer gate are excluded: the full-page
        rasterization already covers whatever is drawn on them. Each paragraph
        keeps the asset it came from so blocks can cite the picture directly.
        """
        if not pages:
            return {}, ()
        candidates: list[tuple[int, Any]] = []
        for page_number in pages:
            if len(candidates) >= _EMBEDDED_MAX_IMAGES:
                logger.warning(
                    "Embedded-image OCR budget reached for %s; capping at %d images",
                    label,
                    _EMBEDDED_MAX_IMAGES,
                )
                break
            try:
                images = list(reader.pages[page_number - 1].images)
            except Exception as exc:  # noqa: BLE001 - best-effort enumeration
                logger.warning(
                    "Unable to enumerate embedded images on page %d of %s: %s",
                    page_number,
                    label,
                    exc,
                )
                continue
            remaining = _EMBEDDED_MAX_IMAGES - len(candidates)
            candidates.extend((page_number, image) for image in images[:remaining])
        embedded: dict[int, list[tuple[str, ImageAsset | None]]] = {}
        assets: list[ImageAsset] = []
        seen_names: set[str] = set()
        streak = 0
        for page_number, image in candidates:
            try:
                paragraphs, asset = self._embedded_image_paragraphs(
                    image, page_number, label, seen_names
                )
            except Exception as exc:  # noqa: BLE001 - best-effort with a streak cap
                streak += 1
                logger.warning(
                    "Skipping embedded image on page %d of %s: %s", page_number, label, exc
                )
                if streak >= _EMBEDDED_PROVIDER_ERROR_LIMIT:
                    _fail_embedded_streak(label, streak, exc)
                continue
            streak = 0
            if paragraphs:
                embedded[page_number] = (
                    embedded.get(page_number, []) + [(paragraph, asset) for paragraph in paragraphs]
                )
        for page_rows in embedded.values():
            for _paragraph, asset in page_rows:
                if asset is not None and not any(
                    existing.name == asset.name for existing in assets
                ):
                    assets.append(asset)
        return embedded, tuple(assets)

    def _embedded_image_paragraphs(
        self, image: Any, page_number: int, label: str, seen_names: set[str]
    ) -> tuple[tuple[str, ...], ImageAsset | None]:
        """Data-quality guards; provider errors propagate to the streak counter.

        Returns the OCR paragraphs plus the image asset they belong to. Asset
        names are unique per document: repeated pypdf names are disambiguated
        with the page number so storage keys stay collision-free.
        """
        try:
            blob = image.data
        except Exception as exc:  # noqa: BLE001 - best-effort per image
            logger.warning(
                "Skipping undecodable embedded image %s in %s: %s", image.name, label, exc
            )
            return (), None
        try:
            from PIL import Image as PILImage
        except ImportError:  # pragma: no cover - dependency ships with the service
            return (), None
        try:
            with PILImage.open(BytesIO(blob)) as frame:
                width, height = frame.size
        except Exception:  # noqa: BLE001 - best-effort per image
            logger.warning("Skipping unreadable embedded image in %s page %d", label, page_number)
            return (), None
        if width < _EMBEDDED_MIN_PIXELS or height < _EMBEDDED_MIN_PIXELS:
            return (), None
        asset_name = str(image.name) if image.name else f"page-{page_number}"
        if asset_name in seen_names:
            asset_name = f"page-{page_number}-{asset_name}"
        seen_names.add(asset_name)
        asset = ImageAsset(
            name=asset_name,
            media_type=_IMAGE_MIME_BY_FORMAT.get((frame.format or "").upper(), ""),
            width=int(width),
            height=int(height),
            blob=blob,
        )
        assert self._ocr is not None  # caller guarantees a service is set
        recognition = self._ocr.ocr_embedded_image_best_effort(
            blob, content_type="", label=f"{label}#page-{page_number}"
        )
        if recognition is None:
            return (), None
        return _ocr_paragraphs(recognition.text, _EMBEDDED_PARAGRAPH_LIMIT), asset


def _visible_chars(paragraphs: tuple[str, ...]) -> int:
    return len("".join(paragraphs).strip())


def _text_source(
    native: dict[int, tuple[str, ...]], ocr_pages: list[int], page_count: int
) -> TextSource:
    if not ocr_pages:
        return "text"
    if len(ocr_pages) == page_count and not any(
        paragraphs for paragraphs in native.values()
    ):
        return "ocr"
    return "mixed"


class ImageParser(ParserAdapter):
    """Standalone image ingestion (.png/.jpg/.jpeg/.webp) through OCR."""

    extensions = (".png", ".jpg", ".jpeg", ".webp")
    content_types = ("image/png", "image/jpeg", "image/webp")

    def __init__(self, ocr: OcrIngestionService):
        self._ocr = ocr

    def parse(
        self, data: bytes, *, file_name: str, content_type: str = ""
    ) -> ParsedDocument:
        recognition = self._ocr.ocr_image(
            data, content_type=content_type or "", label=file_name
        )
        paragraphs = [
            paragraph.strip()
            for paragraph in re.split(r"\n\s*\n", recognition.text)
            if paragraph.strip()
        ]
        if not paragraphs:
            raise PermanentIngestionFailure("Image document produced no text through OCR")
        # The whole file is the image: bind every block to it as a citable
        # asset (loose min-pixel guard — the document itself is the raster).
        asset = _image_asset("source", data, min_pixels=1)
        block_metadata = _image_block_metadata(asset) if asset is not None else {}
        blocks = tuple(
            _block(
                file_name,
                index,
                "page",
                paragraph,
                page_number=1,
                parser_version="ocr-v1",
                metadata=block_metadata,
            )
            for index, paragraph in enumerate(paragraphs)
        )
        return ParsedDocument(
            "document",
            blocks,
            text_source="ocr",
            page_count=1,
            media=(asset,) if asset is not None else (),
        )


def _transcript_blocks(
    file_name: str,
    transcript: TranscriptResult,
    asr: AsrIngestionService,
    media_kind: str,
    *,
    start_ordinal: int = 0,
) -> list[KnowledgeBlock]:
    """Turn transcript segments into retrievable blocks with time provenance.

    Each segment keeps its wall-clock window in block metadata while the
    ``page_number`` anchor maps to the configured timestamp bucket, so the
    existing citation machinery locates a position inside the media without a
    full end-to-end schema change.
    """
    blocks: list[KnowledgeBlock] = []
    for segment in transcript.segments:
        text = segment.text.strip()
        if not text:
            continue
        blocks.append(
            _block(
                file_name,
                start_ordinal + len(blocks),
                "paragraph",
                text,
                page_number=asr.bucket_for_seconds(segment.start_seconds),
                source_start=max(0, int(round(segment.start_seconds))),
                source_end=max(0, int(round(segment.end_seconds))),
                parser_version="asr-v1",
                metadata={
                    "media_kind": media_kind,
                    "start_seconds": segment.start_seconds,
                    "end_seconds": segment.end_seconds,
                },
            )
        )
    return blocks


class AudioParser(ParserAdapter):
    """Standalone audio ingestion (.mp3/.wav/.m4a/.ogg/.opus/.flac) through ASR."""

    extensions = (".mp3", ".wav", ".m4a", ".oga", ".ogg", ".opus", ".flac")
    content_types = (
        "audio/mpeg",
        "audio/wav",
        "audio/x-wav",
        "audio/mp4",
        "audio/x-m4a",
        "audio/ogg",
        "audio/vorbis",
        "audio/opus",
        "audio/flac",
        "audio/x-flac",
    )

    def __init__(self, asr: AsrIngestionService):
        self._asr = asr

    def parse(
        self, data: bytes, *, file_name: str, content_type: str = ""
    ) -> ParsedDocument:
        del content_type  # the registry already matched the extension
        transcript = self._asr.transcribe_audio(data, label=file_name, media_kind="audio")
        blocks = _transcript_blocks(file_name, transcript, self._asr, "audio")
        if not blocks:
            raise PermanentIngestionFailure("Audio document produced no speech transcript")
        return ParsedDocument(
            "document",
            tuple(blocks),
            text_source="asr",
            page_count=blocks[-1].page_number,
        )


class VideoParser(ParserAdapter):
    """Video ingestion (mp4/mkv/mov/webm): audio ASR plus optional keyframe OCR.

    The audio track transcribes exactly like an audio document. When the
    operator turns on keyframe OCR (requires the ingestion OCR stage), sampled
    frames are recognized as extra paragraph blocks carrying their frame
    timestamp and the ``video-keyframe`` OCR source; keyframes themselves are
    not persisted, only their recognized text.
    """

    extensions = (".mp4", ".mkv", ".mov", ".webm")
    content_types = ("video/mp4", "video/x-matroska", "video/quicktime", "video/webm")

    def __init__(
        self,
        asr: AsrIngestionService,
        ocr: OcrIngestionService | None,
        keyframes_ocr: bool,
    ):
        self._asr = asr
        self._ocr = ocr
        self._keyframes_ocr = keyframes_ocr

    def parse(
        self, data: bytes, *, file_name: str, content_type: str = ""
    ) -> ParsedDocument:
        del content_type  # the registry already matched the extension
        transcript = self._asr.transcribe_audio(data, label=file_name, media_kind="video")
        blocks = _transcript_blocks(file_name, transcript, self._asr, "video")
        if not blocks:
            raise PermanentIngestionFailure("Video document produced no speech transcript")
        text_source: TextSource = "asr"
        keyframe_blocks: list[KnowledgeBlock] = []
        if self._keyframes_ocr and self._ocr is not None:
            keyframe_blocks = self._keyframe_ocr_blocks(
                data, file_name, start_ordinal=len(blocks)
            )
            if keyframe_blocks:
                blocks.extend(keyframe_blocks)
                text_source = "mixed"
        return ParsedDocument(
            "document",
            tuple(blocks),
            text_source=text_source,
            page_count=blocks[-1].page_number,
        )

    def _keyframe_ocr_blocks(
        self, data: bytes, file_name: str, *, start_ordinal: int = 0
    ) -> list[KnowledgeBlock]:
        """OCR sampled frames best-effort; provider streaks fail the document."""
        assert self._ocr is not None
        frames = self._asr.extract_video_keyframes(data, label=file_name)
        blocks: list[KnowledgeBlock] = []
        streak = 0
        for seconds, blob in frames:
            try:
                recognition = self._ocr.ocr_embedded_image_best_effort(
                    blob, content_type="", label=f"{file_name}#frame-{seconds:.1f}s"
                )
            except Exception as exc:  # noqa: BLE001 - best-effort with a streak cap
                streak += 1
                logger.warning(
                    "Skipping keyframe at %.1fs in %s: %s", seconds, file_name, exc
                )
                if streak >= _EMBEDDED_PROVIDER_ERROR_LIMIT:
                    _fail_embedded_streak(file_name, streak, exc)
                continue
            streak = 0
            if recognition is None:
                continue
            metadata: dict[str, Any] = {
                "ocr_source": "video-keyframe",
                "keyframe_seconds": round(seconds, 3),
            }
            for paragraph in _ocr_paragraphs(recognition.text, _EMBEDDED_PARAGRAPH_LIMIT):
                blocks.append(
                    _block(
                        file_name,
                        start_ordinal + len(blocks),
                        "paragraph",
                        paragraph,
                        page_number=self._asr.bucket_for_seconds(seconds),
                        source_start=max(0, int(round(seconds))),
                        source_end=max(0, int(round(seconds))),
                        parser_version="ocr-v1",
                        metadata=metadata,
                    )
                )
        return blocks


class DocxParser(ParserAdapter):
    extensions = (".docx",)
    content_types = ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",)

    def __init__(self, ocr: OcrIngestionService | None = None):
        self._ocr = ocr

    def parse(
        self, data: bytes, *, file_name: str, content_type: str = ""
    ) -> ParsedDocument:
        try:
            from docx import Document

            document = Document(BytesIO(data))
        except ImportError as exc:
            raise PermanentIngestionFailure("DOCX parser dependency is not installed") from exc
        except Exception as exc:
            raise PermanentIngestionFailure("DOCX document could not be parsed") from exc
        headings: list[str] = []
        anchors: dict[str, tuple[int, tuple[str, ...], int]] = {}
        pages_seen = 0
        blocks: list[KnowledgeBlock] = []
        for paragraph in document.paragraphs:
            page = pages_seen + 1
            for rid in paragraph._p.xpath(".//a:blip/@r:embed"):
                name = _docx_inline_media_name(document, str(rid))
                if name:
                    anchors[name] = (page, tuple(headings), len(blocks))
            pages_seen += len(
                paragraph._p.xpath('.//w:lastRenderedPageBreak | .//w:br[@w:type="page"]')
            )
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style is not None else "").lower()
            if style.startswith("heading"):
                match = re.search(r"(\d+)", style)
                level = int(match.group(1)) if match else 1
                headings[level - 1 :] = [text]
                block_type: BlockType = "heading"
            elif "list" in style:
                block_type = "list"
            else:
                block_type = "paragraph"
            blocks.append(
                _block(
                    file_name,
                    len(blocks),
                    block_type,
                    text,
                    section_path=tuple(headings),
                    heading_context=headings[-1] if headings else None,
                )
            )
        for table_index, table in enumerate(document.tables, start=1):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            text = "\n".join(" | ".join(row) for row in rows if any(row))
            if text:
                blocks.append(
                    _block(
                        file_name,
                        len(blocks),
                        "table",
                        text,
                        section_path=tuple(headings),
                        table_id=f"table-{table_index}",
                    )
                )
        native_count = len(blocks)
        media_plan: dict[str, list[KnowledgeBlock]] = {}
        media_assets: tuple[ImageAsset, ...] = ()
        if self._ocr is not None and self._ocr.embedded_images:
            media_batches, media_assets = _embedded_ocr_batches(
                self._ocr,
                _media_entries(data, ("word/media/",), file_name),
                file_name,
                "docx",
                start=len(blocks),
                page_number=None,
                anchors=anchors,
            )
            media_plan = dict(media_batches)
        blocks = _interleave_media_blocks(blocks, anchors, media_plan)
        if not blocks:
            raise PermanentIngestionFailure("Document contains no extractable text")
        text_source: TextSource = "text"
        if len(blocks) > native_count:
            text_source = "mixed" if native_count else "ocr"
        return ParsedDocument(
            "document", tuple(blocks), text_source=text_source, media=media_assets
        )


class MarkdownParser(ParserAdapter):
    extensions = (".md", ".markdown")
    content_types = ("text/markdown", "text/x-markdown")

    def parse(
        self, data: bytes, *, file_name: str, content_type: str = ""
    ) -> ParsedDocument:
        text = _decode_utf8(data, "Markdown")
        blocks: list[KnowledgeBlock] = []
        headings: list[str] = []
        current: list[str] = []
        current_type: BlockType = "paragraph"
        in_code = False

        def flush() -> None:
            nonlocal current
            value = "\n".join(current).strip()
            if value:
                blocks.append(
                    _block(
                        file_name,
                        len(blocks),
                        current_type,
                        value,
                        section_path=tuple(headings),
                        heading_context=headings[-1] if headings else None,
                    )
                )
            current = []

        for line in text.splitlines():
            if line.lstrip().startswith("```"):
                flush()
                in_code = not in_code
                current_type = "code" if in_code else "paragraph"
                continue
            if in_code:
                current.append(line)
                continue
            heading = re.match(r"^(#{1,6})\s+(.+)$", line)
            if heading:
                flush()
                level, title = len(heading.group(1)), heading.group(2).strip()
                headings[level - 1 :] = [title]
                blocks.append(
                    _block(
                        file_name,
                        len(blocks),
                        "heading",
                        title,
                        section_path=tuple(headings),
                        heading_context=title,
                    )
                )
                current_type = "paragraph"
            elif re.match(r"^\s*(?:[-*+] |\d+[.)] )", line):
                if current_type != "list":
                    flush()
                    current_type = "list"
                current.append(line.strip())
            elif line.lstrip().startswith(">"):
                if current_type != "quote":
                    flush()
                    current_type = "quote"
                current.append(line.lstrip()[1:].strip())
            elif "|" in line and line.strip().startswith("|"):
                if current_type != "table":
                    flush()
                    current_type = "table"
                current.append(line.strip())
            elif not line.strip():
                flush()
                current_type = "paragraph"
            else:
                if current_type != "paragraph":
                    flush()
                    current_type = "paragraph"
                current.append(line)
        flush()
        if not blocks:
            raise PermanentIngestionFailure("Document contains no extractable text")
        return ParsedDocument("document", tuple(blocks))


class _SafeHtmlParser(HTMLParser):
    ignored = {"script", "style", "noscript", "iframe", "object", "embed"}
    block_tags = {
        "p": "paragraph",
        "li": "list",
        "blockquote": "quote",
        "pre": "code",
        "table": "table",
        **{f"h{i}": "heading" for i in range(1, 7)},
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.depth = 0
        self.active: BlockType | None = None
        self.parts: list[str] = []
        self.items: list[tuple[BlockType, str, int | None]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        tag = tag.lower()
        if tag in self.ignored:
            self.depth += 1
        if self.depth:
            return
        if tag in self.block_tags:
            self._flush()
            self.active = self.block_tags[tag]  # type: ignore[assignment]
            self.heading_level = int(tag[1]) if tag.startswith("h") else None
        elif tag in {"br", "tr"}:
            self.parts.append("\n")
        elif tag in {"td", "th"}:
            self.parts.append(" | ")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in self.ignored and self.depth:
            self.depth -= 1
            return
        if not self.depth and tag in self.block_tags:
            self._flush()

    def handle_data(self, data: str) -> None:
        if not self.depth and self.active is not None:
            self.parts.append(data)

    def close(self) -> None:
        super().close()
        self._flush()

    def _flush(self) -> None:
        text = " ".join("".join(self.parts).split())
        if self.active and text:
            self.items.append((self.active, text, getattr(self, "heading_level", None)))
        self.parts = []
        self.active = None


class HtmlDocumentParser(ParserAdapter):
    extensions = (".html", ".htm")
    content_types = ("text/html", "application/xhtml+xml")

    def parse(
        self, data: bytes, *, file_name: str, content_type: str = ""
    ) -> ParsedDocument:
        text = _decode_utf8(data, "HTML")
        parser = _SafeHtmlParser()
        try:
            parser.feed(text)
            parser.close()
        except Exception as exc:
            raise PermanentIngestionFailure("HTML document could not be parsed") from exc
        headings: list[str] = []
        blocks: list[KnowledgeBlock] = []
        for block_type, value, level in parser.items:
            if block_type == "heading" and level:
                headings[level - 1 :] = [value]
            blocks.append(
                _block(
                    file_name,
                    len(blocks),
                    block_type,
                    value,
                    section_path=tuple(headings),
                    heading_context=headings[-1] if headings else None,
                )
            )
        if not blocks:
            raise PermanentIngestionFailure("Document contains no extractable text")
        return ParsedDocument("document", tuple(blocks))


def _column_name(value: Any, index: int, seen: dict[str, int]) -> str:
    base = re.sub(r"\s+", "_", str(value).strip()).strip("_").lower() if value is not None else ""
    base = base or f"column_{index + 1}"
    seen[base] = seen.get(base, 0) + 1
    return base if seen[base] == 1 else f"{base}_{seen[base]}"


def _infer_type(values: list[Any]) -> str:
    present = [value for value in values if value is not None and value != ""]
    if not present:
        return "string"
    checks = (
        ("boolean", lambda value: isinstance(value, bool)),
        ("integer", lambda value: isinstance(value, int) and not isinstance(value, bool)),
        (
            "decimal",
            lambda value: isinstance(value, (int, float, Decimal)) and not isinstance(value, bool),
        ),
        ("datetime", lambda value: isinstance(value, datetime)),
        ("date", lambda value: isinstance(value, date)),
    )
    for name, check in checks:
        if all(check(value) for value in present):
            return name
    return "string"


def _excel_column(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _tables_from_rows(
    file_name: str, sheet_name: str, rows: list[list[Any]]
) -> tuple[NormalizedTable, ...]:
    row_regions: list[tuple[int, list[list[Any]]]] = []
    current: list[list[Any]] = []
    start = 1
    for row_number, row in enumerate(rows, start=1):
        if any(value is not None and str(value).strip() for value in row):
            if not current:
                start = row_number
            current.append(row)
        elif current:
            row_regions.append((start, current))
            current = []
    if current:
        row_regions.append((start, current))

    regions: list[tuple[int, int, list[list[Any]]]] = []
    for start_row, row_region in row_regions:
        width = max(len(row) for row in row_region)
        active = [
            any(
                index < len(row) and row[index] is not None and str(row[index]).strip()
                for row in row_region
            )
            for index in range(width)
        ]
        band_start: int | None = None
        for index, has_value in enumerate([*active, False]):
            if has_value and band_start is None:
                band_start = index
            elif not has_value and band_start is not None:
                regions.append(
                    (start_row, band_start + 1, [row[band_start:index] for row in row_region])
                )
                band_start = None

    tables: list[NormalizedTable] = []
    for table_index, (start_row, start_column, region) in enumerate(regions, start=1):
        width = max(len(row) for row in region)
        if width > 2_000:
            raise PermanentIngestionFailure("Spreadsheet exceeds the column safety limit")
        padded = [row + [None] * (width - len(row)) for row in region]
        seen: dict[str, int] = {}
        names = [_column_name(value, index, seen) for index, value in enumerate(padded[0])]
        data_rows = padded[1:]
        columns = tuple(
            NormalizedColumn(name, _infer_type([row[index] for row in data_rows]))  # type: ignore[arg-type]
            for index, name in enumerate(names)
        )
        normalized_rows = tuple(
            NormalizedRow(
                row_number=start_row + offset,
                cell_range=(
                    f"{_excel_column(start_column)}{start_row + offset}:"
                    f"{_excel_column(start_column + width - 1)}{start_row + offset}"
                ),
                values={name: row[index] for index, name in enumerate(names)},
            )
            for offset, row in enumerate(data_rows, start=1)
            if any(value is not None and str(value).strip() for value in row)
        )
        if not normalized_rows:
            continue
        if len(normalized_rows) > 250_000:
            raise PermanentIngestionFailure("Spreadsheet exceeds the row safety limit")
        table_id = hashlib.sha256(
            f"{file_name}:{sheet_name}:{start_row}:{table_index}".encode()
        ).hexdigest()[:24]
        tables.append(
            NormalizedTable(
                table_id=table_id,
                sheet_name=sheet_name,
                cell_range=(
                    f"{_excel_column(start_column)}{start_row}:"
                    f"{_excel_column(start_column + width - 1)}{start_row + len(region) - 1}"
                ),
                columns=columns,
                rows=normalized_rows,
            )
        )
    return tuple(tables)


def _spreadsheet_document(
    file_name: str,
    tables: tuple[NormalizedTable, ...],
    *,
    ocr: OcrIngestionService | None = None,
    data: bytes = b"",
    media_sheets: dict[str, str] | None = None,
) -> ParsedDocument:
    blocks: list[KnowledgeBlock] = []
    for table in tables:
        summary = f"Sheet: {table.sheet_name}; Range: {table.cell_range}; Columns: " + ", ".join(
            f"{column.name} ({column.inferred_type})" for column in table.columns
        )
        blocks.append(
            _block(
                file_name,
                len(blocks),
                "table",
                summary,
                sheet_name=table.sheet_name,
                cell_range=table.cell_range,
                table_id=table.table_id,
                metadata={"record_kind": "table_summary"},
            )
        )
        for row in table.rows:
            text = f"Sheet: {table.sheet_name}; Range: {row.cell_range}; " + "; ".join(
                f"{name}: {value}" for name, value in row.values.items() if value not in (None, "")
            )
            blocks.append(
                _block(
                    file_name,
                    len(blocks),
                    "row",
                    text,
                    sheet_name=table.sheet_name,
                    cell_range=row.cell_range,
                    table_id=table.table_id,
                    metadata={"record_kind": "row"},
                )
            )
    if ocr is not None and ocr.embedded_images and data:
        table_block_count = len(blocks)
        media_batches, media_assets = _embedded_ocr_batches(
            ocr,
            _media_entries(data, ("xl/media/",), file_name),
            file_name,
            "xlsx",
            start=len(blocks),
            page_number=None,
            sheets=media_sheets,
        )
        extra = [
            block
            for _name, grouped in media_batches
            for block in grouped
        ]
        blocks.extend(extra)
    else:
        table_block_count = len(blocks)
        media_assets = ()
    text_source: TextSource = "text"
    if len(blocks) > table_block_count:
        text_source = "mixed" if table_block_count else "ocr"
    if not blocks:
        raise PermanentIngestionFailure("Spreadsheet contains no logical tables")
    return ParsedDocument(
        "spreadsheet", tuple(blocks), tables, text_source=text_source, media=media_assets
    )


class CsvParser(ParserAdapter):
    extensions = (".csv",)
    content_types = ("text/csv", "application/csv", "application/vnd.ms-excel")

    def parse(
        self, data: bytes, *, file_name: str, content_type: str = ""
    ) -> ParsedDocument:
        text = _decode_utf8(data, "CSV")
        try:
            sample = "\n".join(text[:8192].splitlines()[:50])
            first_region = re.split(r"\n\s*\n", sample, maxsplit=1)[0]
            dialect = csv.Sniffer().sniff(first_region, delimiters=",;\t")
            rows = [
                [_parse_csv_value(value) for value in row]
                for row in csv.reader(StringIO(text), dialect)
            ]
        except csv.Error as exc:
            raise PermanentIngestionFailure("CSV document could not be parsed") from exc
        tables = _tables_from_rows(file_name, PurePosixPath(file_name).stem, rows)
        return _spreadsheet_document(file_name, tables)


class XlsxParser(ParserAdapter):
    extensions = (".xlsx",)
    content_types = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",)

    def __init__(self, ocr: OcrIngestionService | None = None):
        self._ocr = ocr

    def parse(
        self, data: bytes, *, file_name: str, content_type: str = ""
    ) -> ParsedDocument:
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
        except ImportError as exc:
            raise PermanentIngestionFailure("XLSX parser dependency is not installed") from exc
        except Exception as exc:
            raise PermanentIngestionFailure("XLSX workbook could not be parsed") from exc
        tables: list[NormalizedTable] = []
        try:
            for sheet in workbook.worksheets:
                if sheet.sheet_state != "visible":
                    continue
                rows: list[list[Any]] = []
                formulas: dict[tuple[int, int], str] = {}
                for cells in sheet.iter_rows():
                    row: list[Any] = []
                    for cell in cells:
                        if cell.data_type == "f":
                            formulas[(cell.row, cell.column)] = str(cell.value)
                            row.append(None)
                        else:
                            row.append(cell.value)
                    rows.append(row)
                sheet_tables = _tables_from_rows(file_name, sheet.title, rows)
                for table in sheet_tables:
                    start_column = _column_index(re.split(r"\d", table.cell_range, maxsplit=1)[0])
                    updated_rows: list[NormalizedRow] = []
                    for normalized_row in table.rows:
                        expressions = {
                            column.name: formulas[(normalized_row.row_number, column_index)]
                            for column_index, column in enumerate(table.columns, start=start_column)
                            if (normalized_row.row_number, column_index) in formulas
                        }
                        updated_rows.append(
                            replace(
                                normalized_row,
                                formula_columns=tuple(expressions),
                                formula_expressions=expressions,
                            )
                        )
                    tables.append(replace(table, rows=tuple(updated_rows)))
        finally:
            workbook.close()
        media_sheets: dict[str, str] = {}
        if self._ocr is not None and self._ocr.embedded_images:
            media_sheets = _xlsx_media_sheets(data, file_name)
        return _spreadsheet_document(
            file_name, tuple(tables), ocr=self._ocr, data=data, media_sheets=media_sheets
        )


class DocumentTextExtractor:
    """Parser registry/factory facade for all supported digital formats.

    Media parsing is optional: ``AudioParser``/``VideoParser`` join the
    registry only when an ASR service is wired (and by default the video
    parser stays out entirely), mirroring how the image parser rides the
    optional OCR service.
    """

    def __init__(
        self,
        parsers: tuple[ParserAdapter, ...] | None = None,
        *,
        ocr: OcrIngestionService | None = None,
        asr: AsrIngestionService | None = None,
        video_enabled: bool = False,
        video_keyframe_ocr: bool = False,
    ):
        if parsers is not None:
            self._parsers = parsers
        elif ocr is not None or asr is not None:
            selected: list[ParserAdapter] = []
            if asr is not None:
                selected.append(AudioParser(asr))
                if video_enabled:
                    selected.append(VideoParser(asr, ocr, video_keyframe_ocr))
            if ocr is not None:
                selected.append(ImageParser(ocr))
            selected.extend(
                (
                    PdfParser(ocr),
                    DocxParser(ocr),
                    MarkdownParser(),
                    HtmlDocumentParser(),
                    TextParser(),
                    XlsxParser(ocr),
                    CsvParser(),
                )
            )
            self._parsers = tuple(selected)
        else:
            self._parsers = (
                PdfParser(),
                DocxParser(),
                MarkdownParser(),
                HtmlDocumentParser(),
                TextParser(),
                XlsxParser(),
                CsvParser(),
            )

    def parse(self, data: bytes, *, content_type: str, file_name: str) -> ParsedDocument:
        normalized = content_type.lower().split(";", 1)[0].strip()
        suffix = PurePosixPath(file_name.lower()).suffix
        for parser in self._parsers:
            if suffix in parser.extensions:
                allowed = set(parser.content_types) | {
                    "application/octet-stream",
                    "application/zip",
                }
                if normalized not in allowed and not (
                    isinstance(parser, (TextParser, MarkdownParser, CsvParser))
                    and normalized == "text/plain"
                ):
                    raise PermanentIngestionFailure("File extension and content type do not match")
                return parser.parse(data, file_name=file_name, content_type=normalized)
        raise PermanentIngestionFailure(f"Unsupported document content type: {content_type}")

    def extract(self, data: bytes, *, content_type: str, file_name: str) -> list[ExtractedPage]:
        parsed = self.parse(data, content_type=content_type, file_name=file_name)
        grouped: dict[int, list[str]] = {}
        for block in parsed.blocks:
            page = block.page_number or 1
            grouped.setdefault(page, []).append(block.text)
        return [ExtractedPage(page, "\n\n".join(texts)) for page, texts in sorted(grouped.items())]


def _parse_csv_value(value: str) -> Any:
    cleaned = value.strip()
    if cleaned == "":
        return None
    if re.fullmatch(r"-?(?:0|[1-9]\d*)", cleaned) and not (
        cleaned.startswith("0") and len(cleaned) > 1
    ):
        try:
            return int(cleaned)
        except ValueError:
            pass
    if re.fullmatch(r"-?(?:0|[1-9]\d*)\.\d+", cleaned):
        try:
            return Decimal(cleaned)
        except ValueError:
            pass
    try:
        return date.fromisoformat(cleaned)
    except ValueError:
        return cleaned


def _column_index(letters: str) -> int:
    result = 0
    for letter in letters:
        result = result * 26 + ord(letter) - 64
    return result
